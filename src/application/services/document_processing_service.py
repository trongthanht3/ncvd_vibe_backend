"""
Document processing service for text extraction and sensitive data scanning.

This module handles the actual document processing including:
- PDF to text extraction using OpenAI Vision
- DOCX to text extraction using python-docx
- Text chunking using LangChain
- Sensitive data detection using OpenAI GPT-4
"""

from typing import List, Dict, Any, Optional
from uuid import UUID
import os
import base64
from pathlib import Path
import asyncio
from io import BytesIO

from pdf2image import convert_from_path
from PIL import Image
from docx import Document as DocxDocument
from langchain.text_splitter import MarkdownTextSplitter
import openai
from openai import AsyncOpenAI

from ...data.models.document import DocumentType
from ...data.models.sensitive_data_detection import SensitiveDataType
from ...core.config import get_settings
from ...core.logging import get_logger
from ...core.exceptions import ValidationError

logger = get_logger(__name__)
settings = get_settings()


class DocumentProcessingService:
    """
    Service for document text extraction and sensitive data detection.

    Handles PDF and DOCX processing with AI-powered text extraction
    and sensitive data scanning.
    """

    def __init__(self):
        self.openai_client = AsyncOpenAI(
            api_key=settings.OPENAI_API_KEY
        )
        self.text_splitter = MarkdownTextSplitter(
            chunk_size=2000,
            chunk_overlap=200,
            length_function=len
        )

        # Sensitive data detection patterns
        self.sensitive_data_prompts = self._initialize_detection_prompts()

    async def extract_text_from_document(
        self,
        file_path: str,
        document_type: DocumentType
    ) -> str:
        """
        Extract text from document based on type.

        Args:
            file_path: Path to the document file
            document_type: Type of document (PDF or DOCX)

        Returns:
            Extracted text in markdown format

        Raises:
            ValidationError: If document processing fails
        """
        try:
            if document_type == DocumentType.PDF:
                return await self._extract_text_from_pdf(file_path)
            elif document_type == DocumentType.DOCX:
                return await self._extract_text_from_docx(file_path)
            else:
                raise ValidationError(
                    f"Unsupported document type: {document_type}")
        except Exception as e:
            logger.error(
                f"Failed to extract text from document: {e}", file_path=file_path)
            raise ValidationError(
                f"Failed to extract text from document: {str(e)}")

    async def get_page_count(
        self,
        file_path: str,
        document_type: DocumentType
    ) -> Optional[int]:
        """
        Get the number of pages in the document.

        Args:
            file_path: Path to the document file
            document_type: Type of document

        Returns:
            Number of pages or None if unable to determine
        """
        try:
            if document_type == DocumentType.PDF:
                # Convert PDF to images to count pages
                images = convert_from_path(file_path)
                return len(images)
            elif document_type == DocumentType.DOCX:
                # DOCX doesn't have a direct page count, estimate based on content
                doc = DocxDocument(file_path)
                # Rough estimation: assume ~500 words per page
                total_words = sum(len(paragraph.text.split())
                                  for paragraph in doc.paragraphs)
                return max(1, total_words // 500)
            else:
                return None
        except Exception as e:
            logger.warning(
                f"Failed to get page count: {e}", file_path=file_path)
            return None

    async def scan_for_sensitive_data(
        self,
        text: str,
        document_id: UUID
    ) -> List[Dict[str, Any]]:
        """
        Scan text for sensitive data using OpenAI GPT-4.

        Args:
            text: Text content to scan
            document_id: ID of the document being scanned

        Returns:
            List of detection dictionaries
        """
        detections = []

        try:
            # Split text into chunks for processing
            chunks = self.text_splitter.split_text(text)

            # Process each chunk
            for chunk_index, chunk in enumerate(chunks):
                chunk_detections = await self._scan_chunk_for_sensitive_data(
                    chunk,
                    document_id,
                    chunk_index
                )
                detections.extend(chunk_detections)

            logger.info(
                "Sensitive data scan completed",
                document_id=str(document_id),
                chunks_processed=len(chunks),
                detections_found=len(detections)
            )

            return detections

        except Exception as e:
            logger.error(
                f"Failed to scan for sensitive data: {e}", document_id=str(document_id))
            raise ValidationError(
                f"Failed to scan for sensitive data: {str(e)}")

    # Private methods

    async def _extract_text_from_pdf(self, file_path: str) -> str:
        """
        Extract text from PDF using OpenAI Vision API.

        Args:
            file_path: Path to PDF file

        Returns:
            Extracted text in markdown format
        """
        logger.info("Starting PDF text extraction", file_path=file_path)

        # Convert PDF pages to images
        images = convert_from_path(file_path, dpi=300)

        extracted_pages = []

        # Process each page
        for page_num, image in enumerate(images, 1):
            try:
                # Convert image to base64
                img_buffer = BytesIO()
                image.save(img_buffer, format='PNG')
                img_base64 = base64.b64encode(
                    img_buffer.getvalue()).decode('utf-8')

                # Extract text using OpenAI Vision
                page_text = await self._extract_text_from_image(img_base64, page_num)
                extracted_pages.append(f"## Page {page_num}\n\n{page_text}")

                # Add delay to avoid rate limiting
                if page_num < len(images):
                    await asyncio.sleep(1)

            except Exception as e:
                logger.warning(
                    f"Failed to extract text from page {page_num}: {e}")
                extracted_pages.append(
                    f"## Page {page_num}\n\n[Text extraction failed for this page]")

        # Combine all pages
        full_text = "\n\n".join(extracted_pages)

        logger.info(
            "PDF text extraction completed",
            file_path=file_path,
            pages_processed=len(images),
            text_length=len(full_text)
        )

        return full_text

    async def _extract_text_from_image(self, image_base64: str, page_number: int) -> str:
        """
        Extract text from image using OpenAI Vision API.

        Args:
            image_base64: Base64 encoded image
            page_number: Page number for context

        Returns:
            Extracted text
        """
        prompt = """
Please extract all text from this image and format it as clean markdown.

Instructions:
- Extract ALL visible text, including headers, body text, footnotes, captions, etc.
- Maintain the logical structure using markdown formatting
- Use appropriate headings (##, ###) for titles and sections
- Preserve lists, tables, and other structural elements
- If there are tables, format them as markdown tables
- Include any visible numbers, dates, names, addresses, or other data
- Do not add any interpretation or summary - just extract the raw text
- If text is unclear or partially obscured, use [unclear] placeholder
"""

        try:
            response = await self.openai_client.chat.completions.create(
                model=settings.OPENAI_MODEL,
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": prompt
                            },
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:image/png;base64,{image_base64}"
                                }
                            }
                        ]
                    }
                ],
                max_tokens=4000,
                temperature=0.1
            )

            return response.choices[0].message.content.strip()

        except Exception as e:
            logger.error(
                f"Failed to extract text from image (page {page_number}): {e}")
            raise

    async def _extract_text_from_docx(self, file_path: str) -> str:
        """
        Extract text from DOCX using python-docx.

        Args:
            file_path: Path to DOCX file

        Returns:
            Extracted text in markdown format
        """
        logger.info("Starting DOCX text extraction", file_path=file_path)

        try:
            doc = DocxDocument(file_path)

            content_parts = []

            # Extract paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    # Try to detect heading styles and format accordingly
                    style_name = paragraph.style.name.lower()
                    if 'heading' in style_name:
                        if 'heading 1' in style_name:
                            content_parts.append(f"# {text}")
                        elif 'heading 2' in style_name:
                            content_parts.append(f"## {text}")
                        elif 'heading 3' in style_name:
                            content_parts.append(f"### {text}")
                        else:
                            content_parts.append(f"#### {text}")
                    else:
                        content_parts.append(text)

            # Extract tables
            for table in doc.tables:
                table_text = self._extract_table_as_markdown(table)
                if table_text:
                    content_parts.append(table_text)

            full_text = "\n\n".join(content_parts)

            logger.info(
                "DOCX text extraction completed",
                file_path=file_path,
                paragraphs_processed=len(doc.paragraphs),
                tables_processed=len(doc.tables),
                text_length=len(full_text)
            )

            return full_text

        except Exception as e:
            logger.error(
                f"Failed to extract text from DOCX: {e}", file_path=file_path)
            raise

    def _extract_table_as_markdown(self, table) -> str:
        """
        Extract table content as markdown table.

        Args:
            table: DOCX table object

        Returns:
            Markdown formatted table
        """
        try:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip().replace('\n', ' ')
                         for cell in row.cells]
                if any(cells):  # Only include non-empty rows
                    rows.append(cells)

            if not rows:
                return ""

            # Format as markdown table
            if len(rows) == 1:
                # Single row table
                return f"| {' | '.join(rows[0])} |"

            # Multi-row table with header
            markdown_lines = []

            # Header row
            header = f"| {' | '.join(rows[0])} |"
            markdown_lines.append(header)

            # Separator row
            separator = f"| {' | '.join(['---'] * len(rows[0]))} |"
            markdown_lines.append(separator)

            # Data rows
            for row in rows[1:]:
                # Pad row to match header length
                while len(row) < len(rows[0]):
                    row.append('')
                data_row = f"| {' | '.join(row[:len(rows[0])])} |"
                markdown_lines.append(data_row)

            return "\n".join(markdown_lines)

        except Exception as e:
            logger.warning(f"Failed to extract table as markdown: {e}")
            return ""

    async def _scan_chunk_for_sensitive_data(
        self,
        text_chunk: str,
        document_id: UUID,
        chunk_index: int
    ) -> List[Dict[str, Any]]:
        """
        Scan a text chunk for sensitive data using OpenAI.

        Args:
            text_chunk: Text chunk to scan
            document_id: Document ID
            chunk_index: Index of the chunk

        Returns:
            List of detection dictionaries
        """
        prompt = self._build_sensitive_data_detection_prompt(text_chunk)

        try:
            response = await self.openai_client.chat.completions.create(
                model=settings.OPENAI_MODEL,
                messages=[
                    {
                        "role": "system",
                        "content": "You are a data privacy expert specializing in detecting sensitive personal information in Vietnamese and English documents. You must identify and classify all sensitive data according to Vietnamese data protection standards."
                    },
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],
                max_tokens=2000,
                temperature=0.1,
                response_format={"type": "json_object"}
            )

            # Parse response
            result = response.choices[0].message.content
            import json
            detection_data = json.loads(result)

            # Convert to detection objects
            detections = []
            for detection in detection_data.get('detections', []):
                detection_obj = {
                    'document_id': document_id,
                    'data_type': self._map_detection_type(detection.get('type')),
                    'detected_text': detection.get('text', ''),
                    'masked_text': self._mask_sensitive_text(
                        detection.get('text', ''),
                        detection.get('type')
                    ),
                    'chunk_index': chunk_index,
                    'context': detection.get('context', ''),
                    'confidence_score': float(detection.get('confidence', 0.5)),
                    'detection_pattern': detection.get('pattern', ''),
                    'ai_model_used': settings.OPENAI_MODEL,
                    'detection_metadata': {
                        'chunk_index': chunk_index,
                        'detection_method': 'openai_gpt',
                        'original_type': detection.get('type')
                    }
                }
                detections.append(detection_obj)

            return detections

        except Exception as e:
            logger.error(
                f"Failed to scan chunk for sensitive data: {e}", chunk_index=chunk_index)
            return []

    def _build_sensitive_data_detection_prompt(self, text: str) -> str:
        """
        Build prompt for sensitive data detection.

        Args:
            text: Text to analyze

        Returns:
            Detection prompt
        """
        return f"""
Analyze the following text and identify ALL sensitive personal information according to Vietnamese data protection standards.

Look for these types of sensitive data:
1. **Personal Identification**: CCCD, CMND, Passport numbers, Driver's license
2. **Contact Information**: Phone numbers, email addresses, home addresses
3. **Financial Information**: Bank account numbers, credit card numbers, ATM card numbers
4. **Government IDs**: Tax ID (MST), Social security numbers (BHXH), Employee ID
5. **Medical Information**: Medical records, health conditions, medical ID numbers
6. **Professional Information**: Company internal IDs, confidential business data
7. **Educational Information**: Student IDs, academic records
8. **Other Personal Data**: Any other information that could identify an individual

For EACH piece of sensitive data found, provide:
- The exact text/number found
- The type of sensitive data
- Surrounding context (20-30 words before and after)
- Confidence level (0.0 to 1.0)
- Detection pattern or reasoning

Text to analyze:
```
{text}
```

Respond in JSON format:
{{
  "detections": [
    {{
      "text": "exact sensitive text found",
      "type": "personal_id|phone_number|email|address|bank_account|credit_card|tax_id|social_security|medical_info|financial_info|government_id|custom",
      "context": "surrounding text context",
      "confidence": 0.95,
      "pattern": "description of why this was detected"
    }}
  ]
}}

If no sensitive data is found, return: {{"detections": []}}
"""

    def _map_detection_type(self, ai_type: str) -> SensitiveDataType:
        """
        Map AI detection type to enum.

        Args:
            ai_type: Type string from AI response

        Returns:
            SensitiveDataType enum value
        """
        mapping = {
            'personal_id': SensitiveDataType.PERSONAL_ID,
            'phone_number': SensitiveDataType.PHONE_NUMBER,
            'email': SensitiveDataType.EMAIL,
            'address': SensitiveDataType.ADDRESS,
            'bank_account': SensitiveDataType.BANK_ACCOUNT,
            'credit_card': SensitiveDataType.CREDIT_CARD,
            'tax_id': SensitiveDataType.TAX_ID,
            'social_security': SensitiveDataType.SOCIAL_SECURITY,
            'medical_info': SensitiveDataType.MEDICAL_INFO,
            'financial_info': SensitiveDataType.FINANCIAL_INFO,
            'government_id': SensitiveDataType.GOVERNMENT_ID
        }
        return mapping.get(ai_type, SensitiveDataType.CUSTOM)

    def _mask_sensitive_text(self, text: str, data_type: str) -> str:
        """
        Create masked version of sensitive text.

        Args:
            text: Original sensitive text
            data_type: Type of sensitive data

        Returns:
            Masked text
        """
        if not text:
            return text

        if data_type in ['personal_id', 'bank_account', 'credit_card']:
            # Show first 2 and last 2 characters
            if len(text) > 4:
                return f"{text[:2]}{'*' * (len(text) - 4)}{text[-2:]}"
            else:
                return "*" * len(text)
        elif data_type == 'email':
            # Mask username part
            if "@" in text:
                username, domain = text.split("@", 1)
                if len(username) > 2:
                    masked_username = f"{username[:2]}{'*' * (len(username) - 2)}"
                else:
                    masked_username = "*" * len(username)
                return f"{masked_username}@{domain}"
            return "*" * len(text)
        elif data_type == 'phone_number':
            # Show first 3 and last 2 digits
            digits_only = ''.join(filter(str.isdigit, text))
            if len(digits_only) > 5:
                masked = f"{digits_only[:3]}{'*' * (len(digits_only) - 5)}{digits_only[-2:]}"
                return text.replace(digits_only, masked)
            return "*" * len(text)
        else:
            # Default masking: show first and last character
            if len(text) > 2:
                return f"{text[0]}{'*' * (len(text) - 2)}{text[-1]}"
            else:
                return "*" * len(text)

    def _initialize_detection_prompts(self) -> Dict[str, str]:
        """
        Initialize detection prompts for different data types.

        Returns:
            Dictionary of detection prompts
        """
        return {
            'vietnamese_id': 'Vietnamese CCCD/CMND patterns (12 digits, 9 digits)',
            'phone_vietnam': 'Vietnamese phone numbers (starting with +84, 0)',
            'email_pattern': 'Email addresses with @ symbol',
            'address_vietnam': 'Vietnamese addresses with province/city names',
            'bank_account': 'Bank account numbers (6-20 digits)',
            'credit_card': 'Credit card numbers (16 digits with spaces/dashes)',
            'tax_id_vietnam': 'Vietnamese tax ID (MST) - 10-13 digits',
            'passport': 'Passport numbers (letters and numbers)',
            'license_plate': 'Vietnamese license plates'
        }
